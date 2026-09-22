#!/usr/bin/env python3
"""把 export-sku-rows.js 算出来的 SKU 行写进淘宝官方导入模板，产出可直接上传的文件。

用法:
    python build-sku-excel.py --item <item.json> --template <SKU模板_xls> --out <输出.xlsx>

模板必须用发布页「批量导入」对话框里下载的那一份，列名和列顺序由平台定义，不要自己造。
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl


def load_template(path):
    """平台给的模板扩展名是 .xls，但文件其实是 xlsx（zip）。
    openpyxl 只看扩展名，所以不是 .xlsx 就先复制成临时 .xlsx 再读。"""
    if path.lower().endswith(".xlsx"):
        return openpyxl.load_workbook(path)

    with open(path, "rb") as handle:
        signature = handle.read(4)
    if signature[:2] != b"PK":
        sys.exit(
            "这个模板不是 xlsx（可能是真的老式 xls）。"
            "用 Excel 另存为 .xlsx 之后再跑。"
        )

    temp_dir = tempfile.mkdtemp(prefix="sku-template-")
    temp_path = os.path.join(temp_dir, "template.xlsx")
    shutil.copyfile(path, temp_path)
    return openpyxl.load_workbook(temp_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--item", required=True, help="商品 item.json")
    parser.add_argument("--template", required=True, help="平台下载的 SKU 导入模板")
    parser.add_argument("--out", required=True, help="输出的 xlsx")
    parser.add_argument("--force", action="store_true", help="预检不通过也继续")
    args = parser.parse_args()

    node = os.environ.get("NODE") or shutil.which("node")
    if not node:
        sys.exit("找不到 node，无法导出 SKU 行")

    here = os.path.dirname(os.path.abspath(__file__))
    exporter = os.path.join(here, "export-sku-rows.js")
    cmd = [node, exporter, "--item", args.item]
    if args.force:
        cmd.append("--force")

    result = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if result.returncode != 0:
        sys.exit("导出 SKU 行失败:\n" + (result.stderr or result.stdout))
    data = json.loads(result.stdout)

    workbook = load_template(args.template)
    title = "AI批量上传"
    worksheet = workbook[title] if title in workbook.sheetnames else workbook.worksheets[0]

    # 前几行可能是填写说明，列头行要自己找：包含「适用手机型号」的那一行。
    header_row = None
    header = []
    for r in range(1, min(worksheet.max_row, 10) + 1):
        values = [(worksheet.cell(r, c).value or "") for c in range(1, worksheet.max_column + 1)]
        texts = [str(v).strip() for v in values if v not in (None, "")]
        if "适用手机型号" in texts:
            header_row = r
            header = values
            break
    if header_row is None:
        sys.exit("模板里没找到列头行（应该有一行包含「适用手机型号」）")

    index = {str(name).strip(): i + 1 for i, name in enumerate(header) if name}
    if "适用手机型号" not in index:
        sys.exit("模板缺少「适用手机型号」列，这个模板不对")

    # 我们算出来的列，模板里没有的就跳过并说明——两个模板版本的列不完全一样。
    writable = [name for name in data["columns"] if name in index]
    skipped = [name for name in data["columns"] if name not in index]
    extra = [str(name) for name in header if name and str(name) not in data["columns"]]

    used = [name for name in writable if any(row[data["columns"].index(name)] != "" for row in data["rows"])]

    for offset, row in enumerate(data["rows"]):
        excel_row = header_row + 1 + offset
        for name in writable:
            worksheet.cell(excel_row, index[name]).value = row[data["columns"].index(name)]

    if not args.out.lower().endswith(".xlsx"):
        args.out = args.out + ".xlsx"
    workbook.save(args.out)

    print("已生成: " + args.out)
    print("SKU 行数: " + str(len(data["rows"])))
    print("列头在第 " + str(header_row) + " 行")
    print("用到的轴: " + ", ".join(data.get("usedAxes") or []))
    print("写入的列: " + ", ".join(used))
    if skipped:
        print("模板里没有、已跳过的列: " + ", ".join(skipped))
    if extra:
        print("模板里有、我们不填的列: " + ", ".join(extra))
    for warning in data.get("warnings") or []:
        print("提示: " + warning)


if __name__ == "__main__":
    main()
