#!/usr/bin/env python3
"""解析「发布相似宝贝」页下载的 SKU 模板，还原商品真实的 SKU 结构。

这个模板的第 1 行是填写说明，第 2 行才是列头，数据从第 3 行开始。
"""

import json
import os
import shutil
import sys
import tempfile

import openpyxl


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else (
        r"D:\AutoTaobao\taobao-case\_listing-work\copy-template\SKU模板_150704_1789911422133.xls"
    )
    out_path = sys.argv[2] if len(sys.argv) > 2 else (
        r"D:\AutoTaobao\taobao-case\_listing-work\copy-template\structure.json"
    )

    tmp = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    shutil.copyfile(src, tmp)
    workbook = openpyxl.load_workbook(tmp, data_only=True)
    sheet = workbook["AI批量上传"] if "AI批量上传" in workbook.sheetnames else workbook.worksheets[0]

    # 前几行是填写说明，列头行要自己找：包含「适用手机型号」的那一行才是列头。
    header_row = None
    header = []
    for r in range(1, min(sheet.max_row, 10) + 1):
        values = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        texts = [str(v).strip() for v in values if v not in (None, "")]
        if any("适用手机型号" == t for t in texts):
            header_row = r
            header = values
            break
    if header_row is None:
        sys.exit("没找到列头行（应该有一行包含「适用手机型号」）")

    index = {str(name).strip(): i for i, name in enumerate(header) if name}

    rows = []
    for r in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if any(v not in (None, "") for v in values):
            rows.append(values)

    models = []
    colors = {}
    prices = set()
    stocks = {}
    live = set()
    codes = set()
    for values in rows:
        model = str(values[index["适用手机型号"]] or "").strip()
        color = str(values[index["颜色分类"]] or "").strip()
        if model and model not in models:
            models.append(model)
        if color:
            colors[color] = colors.get(color, 0) + 1
        prices.add(str(values[index["价格"]]))
        stock = str(values[index["数量"]])
        stocks[stock] = stocks.get(stock, 0) + 1
        live.add(str(values[index["是否上架"]]))
        code = str(values[index["商家编码"]] or "").strip()
        if code:
            codes.add(code)

    result = {
        "source": src,
        "headerRow": header_row,
        "columns": [str(h) for h in header if h],
        "dataRows": len(rows),
        "modelCount": len(models),
        "models": models,
        "colorCounts": colors,
        "prices": sorted(prices),
        "stockDistribution": stocks,
        "liveFlags": sorted(live),
        "existingCodes": len(codes),
    }
    with open(out_path, "w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
    print("已写出 " + out_path)
    print("数据行 " + str(len(rows)) + "，机型 " + str(len(models)) + "，颜色 " + str(len(colors)))


if __name__ == "__main__":
    main()
